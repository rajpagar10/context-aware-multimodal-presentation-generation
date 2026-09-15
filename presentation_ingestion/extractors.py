from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol

from .detection import DetectedInput
from .errors import PipelineError
from .media import MediaTools
from .models import ContentItem, ContentKind, ContentStructure, InputKind, PipelineArtifact, SourceTiming
from .transcription import Transcriber
from .utils import normalize_whitespace, stable_id

LOGGER = logging.getLogger(__name__)


@dataclass
class ExtractionResult:
    items: list[ContentItem]
    metadata: dict[str, Any] = field(default_factory=dict)
    artifacts: list[PipelineArtifact] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class Extractor(Protocol):
    def extract(self, detected: DetectedInput, *, output_dir: Path) -> ExtractionResult: ...


def _office_metadata(props: Any) -> dict[str, Any]:
    keys = ("title", "subject", "author", "keywords", "created", "modified")
    return {key: str(value) for key in keys if (value := getattr(props, key, None)) not in (None, "")}


class PptxExtractor:
    def extract(self, detected: DetectedInput, *, output_dir: Path) -> ExtractionResult:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise PipelineError("PPTX support needs python-pptx. Install requirements.txt.", stage="pptx_extraction", path=str(detected.path), modality="pptx") from exc
        try:
            presentation = Presentation(str(detected.path))
        except Exception as exc:
            raise PipelineError(f"Could not open PPTX: {exc}", stage="pptx_extraction", path=str(detected.path), modality="pptx") from exc

        items: list[ContentItem] = []
        for index, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title
            title = normalize_whitespace(title_shape.text) if title_shape and title_shape.has_text_frame else None
            body: list[str] = []
            bullets: list[str] = []
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False) or shape == title_shape:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    text = normalize_whitespace(paragraph.text)
                    if not text:
                        continue
                    body.append(text)
                    if paragraph.level > 0 or len(shape.text_frame.paragraphs) > 1:
                        bullets.append(text)
            notes = None
            try:
                notes_frame = slide.notes_slide.notes_text_frame
                notes = normalize_whitespace(notes_frame.text) if notes_frame and notes_frame.text else None
            except (AttributeError, ValueError):
                LOGGER.debug("Notes are unavailable for slide %s", index)
            item_text = "\n".join(part for part in [title or "", *body] if part)
            items.append(ContentItem(
                id=stable_id("slide", detected.path, index), kind=ContentKind.SLIDE, index=index, text=item_text,
                structure=ContentStructure(title=title, body=body, bullets=bullets, speaker_notes=notes),
                metadata={"slide_number": index},
            ))
        return ExtractionResult(items=items, metadata={"slide_count": len(items), **_office_metadata(presentation.core_properties)})


class PdfExtractor:
    def extract(self, detected: DetectedInput, *, output_dir: Path) -> ExtractionResult:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise PipelineError("PDF support needs pypdf. Install requirements.txt.", stage="pdf_extraction", path=str(detected.path), modality="pdf") from exc
        try:
            reader = PdfReader(str(detected.path))
        except Exception as exc:
            raise PipelineError(f"Could not open PDF: {exc}", stage="pdf_extraction", path=str(detected.path), modality="pdf") from exc
        items: list[ContentItem] = []
        warnings: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = normalize_whitespace(page.extract_text() or "")
            except Exception as exc:
                text = ""
                warnings.append(f"Page {index}: text extraction failed ({exc}).")
            if not text:
                warnings.append(f"Page {index} has no extractable text; it may be blank or image-only (OCR is not enabled in Review 1).")
            items.append(ContentItem(
                id=stable_id("page", detected.path, index), kind=ContentKind.PAGE, index=index, text=text,
                structure=ContentStructure(title=f"Page {index}", body=[text] if text else []), metadata={"page_number": index},
            ))
        raw_metadata = reader.metadata or {}
        metadata = {str(key).lstrip("/"): str(value) for key, value in raw_metadata.items() if value is not None}
        metadata["page_count"] = len(items)
        return ExtractionResult(items=items, metadata=metadata, warnings=warnings)


class TextExtractor:
    def extract(self, detected: DetectedInput, *, output_dir: Path) -> ExtractionResult:
        text = detected.path.read_text(encoding="utf-8-sig")
        title = detected.path.stem
        return ExtractionResult(items=[ContentItem(
            id=stable_id("text", detected.path, 1), kind=ContentKind.TEXT_SECTION, index=1, text=text,
            structure=ContentStructure(title=title, body=[text]), metadata={"character_count": len(text)},
        )], metadata={"encoding": "utf-8", "character_count": len(text)})


class DocxExtractor:
    def extract(self, detected: DetectedInput, *, output_dir: Path) -> ExtractionResult:
        try:
            from docx import Document
            document = Document(str(detected.path))
        except Exception as exc:
            raise PipelineError(f"Could not open DOCX: {exc}", stage="docx_extraction", path=str(detected.path), modality="docx") from exc
        paragraphs = [normalize_whitespace(paragraph.text) for paragraph in document.paragraphs if normalize_whitespace(paragraph.text)]
        text = "\n".join(paragraphs)
        return ExtractionResult(items=[ContentItem(
            id=stable_id("docx", detected.path, 1), kind=ContentKind.DOCUMENT_SECTION, index=1, text=text,
            structure=ContentStructure(title=detected.path.stem, body=paragraphs), metadata={"paragraph_count": len(paragraphs)},
        )], metadata={"paragraph_count": len(paragraphs), **_office_metadata(document.core_properties)})


class AudioExtractor:
    def __init__(self, transcriber: Transcriber, media_tools: MediaTools | None = None):
        self.transcriber = transcriber
        self.media_tools = media_tools or MediaTools()

    def extract(self, detected: DetectedInput, *, output_dir: Path) -> ExtractionResult:
        try:
            metadata = self.media_tools.probe(detected.path)
            transcript = self.transcriber.transcribe(detected.path)
        except PipelineError as exc:
            _add_error_context(exc, detected)
            raise
        items = _transcript_items(detected, transcript.segments)
        metadata["transcription"] = {"engine": transcript.engine, "language": transcript.language, "duration_seconds": transcript.duration_seconds}
        return ExtractionResult(items=items, metadata=metadata)


class VideoExtractor:
    def __init__(self, transcriber: Transcriber, media_tools: MediaTools | None = None):
        self.transcriber = transcriber
        self.media_tools = media_tools or MediaTools()

    def extract(self, detected: DetectedInput, *, output_dir: Path) -> ExtractionResult:
        try:
            metadata = self.media_tools.probe(detected.path)
            audio_path = output_dir / "derived" / f"{detected.path.stem}_audio_16khz_mono.wav"
            extracted = self.media_tools.extract_audio(detected.path, audio_path)
            transcript = self.transcriber.transcribe(extracted)
        except PipelineError as exc:
            _add_error_context(exc, detected)
            raise
        items = _transcript_items(detected, transcript.segments)
        metadata["transcription"] = {"engine": transcript.engine, "language": transcript.language, "duration_seconds": transcript.duration_seconds}
        return ExtractionResult(items=items, metadata=metadata, artifacts=[PipelineArtifact(kind="extracted_audio", path=str(extracted), metadata={"sample_rate_hz": 16000, "channels": 1})])


def _transcript_items(detected: DetectedInput, segments: list[Any]) -> list[ContentItem]:
    if not segments:
        return [ContentItem(id=stable_id("segment", detected.path, 1), kind=ContentKind.TRANSCRIPT_SEGMENT, index=1, text="", metadata={"transcription_empty": True})]
    return [ContentItem(
        id=stable_id("segment", detected.path, segment.index), kind=ContentKind.TRANSCRIPT_SEGMENT, index=segment.index,
        text=segment.text, structure=ContentStructure(body=[segment.text]),
        source_timing=SourceTiming(start_seconds=segment.start_seconds, end_seconds=segment.end_seconds, provenance="asr"),
        metadata={"word_timestamps": segment.words},
    ) for segment in segments]


def _add_error_context(error: PipelineError, detected: DetectedInput) -> None:
    """Preserve the original stage while adding missing user-facing context."""
    if error.path is None:
        error.path = str(detected.path)
    if error.modality is None:
        error.modality = detected.kind.value


def build_extractor(kind: InputKind, *, transcriber: Transcriber, media_tools: MediaTools | None = None) -> Extractor:
    extractors: dict[InputKind, Extractor] = {
        InputKind.PPTX: PptxExtractor(), InputKind.PDF: PdfExtractor(), InputKind.TEXT: TextExtractor(), InputKind.DOCX: DocxExtractor(),
        InputKind.AUDIO: AudioExtractor(transcriber, media_tools), InputKind.VIDEO: VideoExtractor(transcriber, media_tools),
    }
    return extractors[kind]
