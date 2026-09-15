from __future__ import annotations

import json
import tempfile
import unittest
import wave
from pathlib import Path

from pydantic import TypeAdapter
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas

from presentation_ingestion.detection import detect_input
from presentation_ingestion.errors import PipelineError, UnsupportedInputError
from presentation_ingestion.extractors import VideoExtractor
from presentation_ingestion.instructions import RuleBasedInstructionParser
from presentation_ingestion.models import InputKind, NormalizedRepresentation, PipelineArtifact
from presentation_ingestion.pipeline import IngestionPipeline, PipelineConfig
from presentation_ingestion.transcription import TranscriptSegment, TranscriptionResult


class StubTranscriber:
    def transcribe(self, media_path: Path) -> TranscriptionResult:
        return TranscriptionResult(
            language="en", duration_seconds=2.0, engine="stub_asr",
            segments=[TranscriptSegment(1, "The accuracy is ninety four percent.", 0.0, 2.0, [{"word": "accuracy", "start_seconds": 0.4, "end_seconds": 0.9}])],
        )


class FakeMediaTools:
    def probe(self, path: Path) -> dict:
        return {"format": {"format_name": "mov,mp4,m4a,3gp,3g2,mj2", "duration": "2.0"}, "streams": [{"codec_type": "video"}, {"codec_type": "audio"}]}

    def extract_audio(self, video_path: Path, destination: Path) -> Path:
        destination.parent.mkdir(parents=True, exist_ok=True)
        _write_wav(destination)
        return destination


def _write_wav(path: Path) -> None:
    with wave.open(str(path), "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(16000)
        output.writeframes(b"\x00\x00" * 16000)


def _write_pptx(path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Introduction"
    body = slide.placeholders[1].text_frame
    body.text = "Project context"
    body.add_paragraph().text = "Accuracy is 94 percent"
    slide.notes_slide.notes_text_frame.text = "Introduce the project clearly."
    deck.save(path)


def _write_pdf(path: Path) -> None:
    canvas = Canvas(str(path))
    canvas.drawString(72, 720, "Results: accuracy is 94 percent.")
    canvas.showPage()
    canvas.drawString(72, 720, "Conclusion")
    canvas.save()


class ReviewOnePipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.pptx = self.root / "deck.pptx"
        self.pdf = self.root / "report.pdf"
        self.text = self.root / "script.txt"
        self.wav = self.root / "talk.wav"
        _write_pptx(self.pptx)
        _write_pdf(self.pdf)
        self.text.write_text("Welcome to the introduction. The accuracy is 94 percent.", encoding="utf-8")
        _write_wav(self.wav)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def pipeline(self, name: str, *, generate_narration: bool = False) -> IngestionPipeline:
        return IngestionPipeline(PipelineConfig(self.root / name, generate_narration=generate_narration), transcriber=StubTranscriber())

    def test_pptx_extraction_narration_and_semantic_plan(self) -> None:
        result, output = self.pipeline("pptx", generate_narration=True).run(
            self.pptx, instruction="Speak calmly during the introduction and emphasize the accuracy value."
        )
        self.assertEqual(result.input.kind, InputKind.PPTX)
        self.assertEqual(result.input.metadata["slide_count"], 1)
        self.assertEqual(result.content_items[0].structure.title, "Introduction")
        self.assertIn("Accuracy is 94 percent", result.content_items[0].text)
        self.assertIn("Introduce the project", result.content_items[0].structure.speaker_notes or "")
        self.assertEqual(result.narration[0].origin, "generated")
        self.assertEqual(result.narration[0].content_item_ids, [result.content_items[0].id])
        self.assertEqual(len(result.semantic_control_plan.directives), 2)
        self.assertTrue(all(directive.timing_status == "semantic_pending_alignment" for directive in result.semantic_control_plan.directives))
        self.assertTrue(output.exists())

    def test_pdf_extraction_preserves_page_order_and_generates_mapping(self) -> None:
        result, _ = self.pipeline("pdf", generate_narration=True).run(self.pdf, instruction="Become energetic during results.")
        self.assertEqual(result.input.kind, InputKind.PDF)
        self.assertEqual([item.index for item in result.content_items], [1, 2])
        self.assertIn("accuracy is 94 percent", result.content_items[0].text.casefold())
        self.assertEqual(len(result.narration), 2)
        self.assertEqual(result.narration[0].content_item_ids, [result.content_items[0].id])

    def test_text_is_preserved_as_user_script(self) -> None:
        original = self.text.read_text(encoding="utf-8")
        result, _ = self.pipeline("text").run(self.text, instruction="Pause before the introduction.")
        self.assertEqual(result.input.kind, InputKind.TEXT)
        self.assertEqual(result.content_items[0].text, original)
        self.assertEqual(result.narration[0].origin, "user_script")
        self.assertEqual(result.narration[0].text, original)
        self.assertTrue(result.semantic_control_plan.directives[0].attributes.pause_before)

    def test_wav_audio_has_actual_metadata_and_asr_timestamps(self) -> None:
        result, _ = self.pipeline("audio").run(self.wav)
        self.assertEqual(result.input.kind, InputKind.AUDIO)
        self.assertEqual(result.input.metadata["streams"][0]["sample_rate"], "16000")
        item = result.content_items[0]
        self.assertEqual(item.source_timing.provenance, "asr")
        self.assertEqual((item.source_timing.start_seconds, item.source_timing.end_seconds), (0.0, 2.0))
        self.assertEqual(result.narration[0].origin, "source_transcript")

    def test_video_extractor_creates_derived_audio_and_preserves_asr_time(self) -> None:
        video = self.root / "talk.mp4"
        video.write_bytes(b"not-a-real-video-because-ffmpeg-is-isolated-in-this-unit-test")
        detected = detect_input(video)
        extracted = VideoExtractor(StubTranscriber(), FakeMediaTools()).extract(detected, output_dir=self.root / "video")
        self.assertEqual(extracted.items[0].text, "The accuracy is ninety four percent.")
        self.assertEqual(extracted.items[0].source_timing.provenance, "asr")
        self.assertEqual(extracted.artifacts[0].kind, "extracted_audio")
        self.assertTrue(Path(extracted.artifacts[0].path).exists())

    def test_schema_json_round_trip(self) -> None:
        result, _ = self.pipeline("schema", generate_narration=True).run(self.pptx)
        serialized = json.dumps(result.model_dump(mode="json"))
        restored = NormalizedRepresentation.model_validate_json(serialized)
        self.assertEqual(restored.job_id, result.job_id)
        self.assertEqual(restored.content_items[0].id, result.content_items[0].id)

    def test_instruction_parser_does_not_create_timestamps(self) -> None:
        result, _ = self.pipeline("instructions").run(self.text)
        plan = RuleBasedInstructionParser().parse(
            "Speak calmly during the introduction, become energetic during the results, emphasize the accuracy value, and pause before the conclusion.",
            result.content_items,
        )
        self.assertEqual(len(plan.directives), 4)
        self.assertEqual(plan.directives[0].attributes.emotion, "calm")
        self.assertEqual(plan.directives[1].attributes.energy, 0.8)
        self.assertEqual(plan.directives[2].attributes.emphasis, 0.9)
        self.assertTrue(plan.directives[3].attributes.pause_before)
        self.assertFalse(hasattr(plan.directives[0], "start_time"))

    def test_unsupported_extension_is_actionable(self) -> None:
        invalid = self.root / "notes.csv"
        invalid.write_text("a,b", encoding="utf-8")
        with self.assertRaises(UnsupportedInputError) as raised:
            detect_input(invalid)
        self.assertIn("Unsupported input format", str(raised.exception))

    def test_malformed_pdf_is_rejected_during_validation(self) -> None:
        malformed = self.root / "broken.pdf"
        malformed.write_bytes(b"this is not a PDF")
        with self.assertRaises(PipelineError) as raised:
            detect_input(malformed)
        self.assertIn("PDF signature", str(raised.exception))


if __name__ == "__main__":
    unittest.main()
